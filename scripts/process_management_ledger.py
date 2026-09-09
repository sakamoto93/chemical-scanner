#!/usr/bin/env python3
"""
化学物質管理簿の CAS番号追加・リスクアセスメント対象判定スクリプト

「管理簿」シートの各行は「生工研No.」を持ち、この番号で
「化学物質一覧（書込み不可）」シート（毒物及び劇物取締法の対象物質マスタ）を
引くと CAS番号が分かる。これを利用して：

1. 管理簿の各行に CAS番号を自動転記
2. その CAS番号を、労働安全衛生法リスクアセスメント対象化合物リスト
   （data/risk_assessment.csv）と照合し、対象かどうかを判定
3. 結果を新しい列として管理簿シートに追記した Excel ファイルを出力

生工研No.が数値でない・特定できない行（"735？" "608-619" のような
あいまいな手入力）は自動処理せず、「要確認」として明示的にフラグを立てる
（法令上の該当判断は化学物質管理者の確認が必要なため、スクリプト側で
 推測結果を断定的に埋めることはしない）。

使い方:
    python scripts/process_management_ledger.py <元のExcelファイル> \\
        [--risk-csv data/risk_assessment.csv] \\
        [--output 出力ファイル名.xlsx]
"""
import sys
import argparse
import re
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill


def load_risk_db(risk_csv_path):
    risk_db = {}
    with open(risk_csv_path, newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            cas = row.get("cas_number", "").strip()
            if cas:
                risk_db[cas] = {
                    "name": row.get("compound_name", "").strip(),
                    "related_cas": [c.strip() for c in row.get("related_cas", "").split(",") if c.strip()],
                }
    return risk_db


def build_master_map(ws_master):
    """化学物質一覧（書込み不可）シートから 生工研No. -> {name, cas} の対応表を作る"""
    master_map = {}
    for row_idx, row in enumerate(ws_master.iter_rows(values_only=True), 1):
        if row_idx == 1:
            continue  # ヘッダー行
        no = row[0]
        name = row[1]
        cas_raw = row[14] if len(row) > 14 else None
        if no is None:
            continue
        cas = None
        if cas_raw:
            m = re.search(r'\d{2,7}-\d{2}-\d', str(cas_raw))
            if m:
                cas = m.group(0)
        master_map[no] = {"name": name, "cas": cas}
    return master_map


def process(input_path, risk_csv_path, output_path):
    risk_db = load_risk_db(risk_csv_path)
    print(f"✅ リスクアセスメント対象DB読み込み: {len(risk_db)}件\n")

    wb = openpyxl.load_workbook(input_path, data_only=False)
    ws_master_ro = openpyxl.load_workbook(input_path, data_only=True)['化学物質一覧（書込み不可）']
    master_map = build_master_map(ws_master_ro)
    print(f"✅ 化学物質一覧マスタ読み込み: {len(master_map)}件\n")

    ws_ledger = wb['管理簿']

    HEADER_ROW = 6
    CAS_COL = 11   # K列に新設
    RISK_COL = 12  # L列に新設
    NOTE_COL = 13  # M列に新設（要確認事項）

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    ws_ledger.cell(row=HEADER_ROW, column=CAS_COL, value="CAS番号")
    ws_ledger.cell(row=HEADER_ROW, column=RISK_COL, value="リスクアセスメント対象")
    ws_ledger.cell(row=HEADER_ROW, column=NOTE_COL, value="備考（自動処理）")
    for col in (CAS_COL, RISK_COL, NOTE_COL):
        cell = ws_ledger.cell(row=HEADER_ROW, column=col)
        cell.fill = header_fill
        cell.font = header_font

    stats = {"matched": 0, "risk_target": 0, "unresolved": 0, "no_cas_in_master": 0, "skipped_blank": 0}
    unresolved_rows = []

    for row_idx in range(HEADER_ROW + 1, ws_ledger.max_row + 1):
        seiko_no_cell = ws_ledger.cell(row=row_idx, column=3)
        seiko_no = seiko_no_cell.value
        name_cell = ws_ledger.cell(row=row_idx, column=4)

        if seiko_no is None:
            stats["skipped_blank"] += 1
            continue  # 未登録の空き管理番号行（#N/A行）はスキップ

        if not isinstance(seiko_no, (int, float)):
            # "735？" "608-619" のようなあいまいな手入力 → 要確認としてフラグ
            stats["unresolved"] += 1
            unresolved_rows.append((row_idx, seiko_no, name_cell.value))
            ws_ledger.cell(row=row_idx, column=NOTE_COL,
                            value=f"生工研No.が「{seiko_no}」であいまいなため要確認（化学物質管理者による判定が必要）")
            continue

        info = master_map.get(seiko_no)
        if info is None:
            stats["unresolved"] += 1
            unresolved_rows.append((row_idx, seiko_no, name_cell.value))
            ws_ledger.cell(row=row_idx, column=NOTE_COL, value="生工研No.が化学物質一覧に見つからない")
            continue

        if not info["cas"]:
            stats["no_cas_in_master"] += 1
            ws_ledger.cell(row=row_idx, column=NOTE_COL, value="化学物質一覧にCAS番号の記載なし")
            continue

        cas = info["cas"]
        stats["matched"] += 1
        ws_ledger.cell(row=row_idx, column=CAS_COL, value=cas)

        risk_info = risk_db.get(cas)
        if risk_info:
            stats["risk_target"] += 1
            ws_ledger.cell(row=row_idx, column=RISK_COL, value="対象")
            if len(risk_info["related_cas"]) > 1:
                ws_ledger.cell(row=row_idx, column=NOTE_COL,
                                value=f"関連CAS: {', '.join(risk_info['related_cas'])}")
        else:
            ws_ledger.cell(row=row_idx, column=RISK_COL, value="-")

    wb.save(output_path)

    print("=" * 60)
    print("処理結果サマリー")
    print("=" * 60)
    print(f"CAS番号を自動転記できた行数: {stats['matched']}")
    print(f"  うちリスクアセスメント対象: {stats['risk_target']}")
    print(f"化学物質一覧にCAS記載がなかった行数: {stats['no_cas_in_master']}")
    print(f"要確認（生工研No.があいまい/未登録）: {stats['unresolved']}")
    for r in unresolved_rows:
        print(f"  行{r[0]}: 生工研No.={r[1]!r}, 名称={r[2]}")
    print(f"空き管理番号行（スキップ）: {stats['skipped_blank']}")
    print()
    print(f"💾 出力ファイル: {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx", help="元の化学物質管理簿.xlsxファイル")
    parser.add_argument("--risk-csv", default="data/risk_assessment.csv")
    parser.add_argument("--output", default=None)
    args = parser.parse_args()

    output = args.output or args.input_xlsx.replace(".xlsx", "_処理済み.xlsx")
    process(args.input_xlsx, args.risk_csv, output)
